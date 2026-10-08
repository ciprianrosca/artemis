#!/usr/bin/env python3
"""Turn a CV into plain text for `artemis setup`: PDF, DOCX, TXT or MD, or a public LinkedIn URL.

  python tools/read_cv.py ~/Downloads/my_cv.pdf
  python tools/read_cv.py ~/Downloads/my_cv.docx
  python tools/read_cv.py https://www.linkedin.com/in/<slug>/

A LinkedIn "Save to PDF" export is a normal PDF. Old .doc files can't be read: open them in Word or
Google Docs and save as .docx or .pdf. PDF needs PyMuPDF and DOCX needs python-docx
(`pip install -r requirements.txt`).
"""
import json
import os
import re
import sys

from common import utf8_console

utf8_console()


def read_pdf(path):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    with fitz.open(path) as doc:
        text = "\n".join(page.get_text() for page in doc)
    if len(text.strip()) < 200:
        print("WARN: almost no text found. This may be a scanned image; ask for a text PDF, the DOCX, "
              "or the LinkedIn profile instead.", file=sys.stderr)
    return text


def read_docx(path):
    from docx import Document
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many CV templates put whole sections in tables
        for row in table.rows:
            lines.append(" | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())))
    return "\n".join(lines)


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    if re.match(r"https?://(www\.)?linkedin\.com/in/", src):
        from linkedin_jobs import profile
        print(json.dumps(profile(src), ensure_ascii=False, indent=2))
        return
    path = os.path.expanduser(src)
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        text = read_pdf(path)
    elif ext == ".docx":
        text = read_docx(path)
    elif ext in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="replace") as f:
            text = f.read()
    elif ext == ".doc":
        sys.exit("Old .doc format: open it in Word or Google Docs and save it as .docx or .pdf.")
    else:
        sys.exit(f"Unsupported file type '{ext}'. Use PDF, DOCX, TXT or MD.")
    print(re.sub(r"\n{3,}", "\n\n", text).strip())


if __name__ == "__main__":
    main()
