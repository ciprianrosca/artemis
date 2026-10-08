#!/usr/bin/env python3
"""Render a tailored resume / cover letter from Markdown to ATS-safe DOCX and PDF.

  python tools/render_resume.py applications/2026-10-08_acme_staff-eng/resume.md
  python tools/render_resume.py <file.md> --name Sam_Rivera_CV     # output file name the recruiter sees
  python tools/render_resume.py <file.md> --no-pdf

Markdown subset: `# Name`, `## Section`, `### Role line`, paragraphs, `- bullets`, **bold**.
`{{contact}}` is replaced with the line built from profile/contact.json (gitignored, so the phone
number never lands in the repo).

Layout is deliberately plain: one column, no tables, no text boxes, no images, standard fonts —
that is what ATS parsers read reliably. DOCX needs python-docx; PDF needs PyMuPDF
(`pip install -r requirements.txt`).
"""
import argparse
import html
import json
import os
import re
import sys

from common import profile_path, utf8_console

utf8_console()

CONTACT = profile_path("contact.json")
FONT = "Calibri"


def contact_line():
    try:
        with open(CONTACT, encoding="utf-8") as f:
            c = json.load(f)
    except FileNotFoundError:
        print("WARN: profile/contact.json missing — {{contact}} left empty", file=sys.stderr)
        return ""
    return " · ".join(c[k] for k in ("location", "email", "phone", "linkedin") if c.get(k))


def parse(md):
    md = md.replace("{{contact}}", contact_line())
    blocks = []
    for line in md.splitlines():
        s = line.rstrip()
        if not s.strip():
            blocks.append(("blank", ""))
        elif s.startswith("### "):
            blocks.append(("h3", s[4:]))
        elif s.startswith("## "):
            blocks.append(("h2", s[3:]))
        elif s.startswith("# "):
            blocks.append(("h1", s[2:]))
        elif re.match(r"\s*[-*] ", s):
            blocks.append(("li", re.sub(r"^\s*[-*] ", "", s)))
        elif s.startswith("<!--"):
            continue
        else:
            blocks.append(("p", s.strip()))
    return blocks


def runs(text):
    """Split **bold** segments -> [(text, bold)]."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") else (p, False) for p in parts if p]


def to_docx(blocks, out):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.9)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = FONT, Pt(10.5)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    base.paragraph_format.space_after = Pt(2)

    def para(text, size=10.5, bold=False, align=None, before=0, after=2, color=None, style=None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if align is not None:
            p.alignment = align
        for t, b in runs(text):
            r = p.add_run(t)
            r.bold = bold or b
            r.font.size = Pt(size)
            r.font.name = FONT
            if color:
                r.font.color.rgb = RGBColor(*color)
        return p

    def rule(p):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "888888")):
            bottom.set(qn(f"w:{k}"), v)
        bdr.append(bottom)
        pPr.append(bdr)

    seen_h2 = False
    for kind, text in blocks:
        if kind == "h1":
            para(text, 20, True, WD_ALIGN_PARAGRAPH.CENTER, after=0)
        elif kind == "h2":
            seen_h2 = True
            rule(para(text.upper(), 11, True, before=8, after=3, color=(0x22, 0x3A, 0x5E)))
        elif kind == "h3":
            para(text, 10.5, True, before=5, after=1)
        elif kind == "li":
            p = para(text, style="List Bullet", after=1)
            p.paragraph_format.left_indent = Cm(0.6)
        elif kind == "p":
            header = not seen_h2
            para(text, 10 if header else 10.5, align=WD_ALIGN_PARAGRAPH.CENTER if header else None)
    doc.save(out)


def to_html(blocks):
    out, in_list, seen_h2 = [], False, False
    for kind, text in blocks:
        t = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", html.escape(text))
        if kind != "li" and in_list:
            out.append("</ul>")
            in_list = False
        if kind == "h1":
            out.append(f"<h1>{t}</h1>")
        elif kind == "h2":
            seen_h2 = True
            out.append(f"<h2>{t}</h2>")
        elif kind == "h3":
            out.append(f"<h3>{t}</h3>")
        elif kind == "li":
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{t}</li>")
        elif kind == "p":
            out.append(f'<p class="{"head" if not seen_h2 else ""}">{t}</p>')
    if in_list:
        out.append("</ul>")
    css = """@page{size:A4;margin:16mm 19mm}body{font-family:Calibri,Carlito,Arial,sans-serif;font-size:10.5pt;
color:#111;line-height:1.3}h1{font-size:20pt;text-align:center;margin:0}p.head{text-align:center;margin:1pt 0;
font-size:10pt}h2{font-size:11pt;text-transform:uppercase;color:#223a5e;border-bottom:1px solid #888;
margin:9pt 0 3pt;padding-bottom:1pt}h3{font-size:10.5pt;margin:6pt 0 1pt}p{margin:0 0 3pt}ul{margin:0 0 3pt;
padding-left:16pt}li{margin:0 0 1pt}"""
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head><body>{''.join(out)}</body></html>"


def to_pdf(blocks, out):
    """HTML → PDF with PyMuPDF's Story layout engine: local, no browser needed."""
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz
        except ImportError:
            print("WARN: PyMuPDF not installed, so no PDF. Run `pip install pymupdf`, or open the DOCX "
                  "in Word or Google Docs and export it as PDF.", file=sys.stderr)
            return False
    body = to_html(blocks)
    css = re.search(r"<style>(.*?)</style>", body, re.S).group(1)
    css = re.sub(r"@page\{[^}]*\}", "", css).replace("Calibri,Carlito,Arial,", "")
    story = fitz.Story(html=re.sub(r"<style>.*?</style>", "", body, flags=re.S), user_css=css)
    page = fitz.paper_rect("a4")
    content = page + (54, 45, -54, -45)  # ~19 mm sides, ~16 mm top/bottom
    writer = fitz.DocumentWriter(out)
    more = 1
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(content)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--no-docx", action="store_true")
    ap.add_argument("--name", help="output file name without extension (default: same as the .md)")
    a = ap.parse_args()
    with open(a.markdown, encoding="utf-8") as f:
        blocks = parse(f.read())
    stem = os.path.splitext(a.markdown)[0]
    if a.name:
        stem = os.path.join(os.path.dirname(a.markdown), a.name)
    if not a.no_docx:
        to_docx(blocks, stem + ".docx")
        print(stem + ".docx")
    if not a.no_pdf and to_pdf(blocks, stem + ".pdf"):
        print(stem + ".pdf")


if __name__ == "__main__":
    main()
